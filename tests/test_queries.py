"""Several search queries in one run: what is collected, what lands in the report, what counts as a rival."""

from __future__ import annotations

import asyncio
from datetime import datetime

import pytest
from openpyxl import load_workbook

from mpparser.competitors import compare_with_rivals, parse_own
from mpparser.export.excel import build_file_name, export_to_excel
from mpparser.models import Product
from mpparser.monitoring import job_key
from mpparser.runner import _once_per_product, _search_all, _share_details
from mpparser.settings import MAX_QUERIES, InputMode, ParseSettings

TWO = "чайник\nтермокружка"


def product(article: str, price: float | None, query: str = "", marketplace: str = "Wildberries",
            brand: str = "") -> Product:
    return Product(marketplace=marketplace, article=article, name=f"Товар {article}", price=price, query=query,
                   brand=brand, url=f"https://wb/{article}")


class FakeParser:
    """A marketplace that answers every query with products of its own."""

    title = "Wildberries"

    def __init__(self, answers: dict[str, list[Product]]) -> None:
        self.answers = answers
        self.asked: list[str] = []

    async def search(self, query: str, limit: int, sort: object) -> list[Product]:
        self.asked.append(query)
        return [product(p.article, p.price) for p in self.answers.get(query, [])]


class Silent:
    def set_span(self, *_: float) -> None: ...

    def log(self, *_: object) -> None: ...


# --- settings ---


def test_every_line_is_a_query():
    assert ParseSettings(query=TWO).queries == ["чайник", "термокружка"]
    assert ParseSettings(query="  чайник  \n\n  ").queries == ["чайник"]
    assert ParseSettings(query="   ").queries == []


def test_the_subject_names_the_first_query_and_counts_the_rest():
    assert ParseSettings(query="чайник").subject() == "чайник"
    assert ParseSettings(query=TWO).subject() == "чайник и ещё 1"
    assert ParseSettings(mode=InputMode.IDS).subject() == "артикулы"


def test_an_empty_box_and_too_many_queries_are_refused():
    assert "Введите поисковый запрос." in ParseSettings(query=" \n ").validate()
    many = ParseSettings(query="\n".join(f"запрос {n}" for n in range(MAX_QUERIES + 1)))
    assert any("не больше" in problem for problem in many.validate())
    assert ParseSettings(query=TWO).validate() == []


def test_a_different_list_of_queries_is_a_different_history():
    assert job_key(ParseSettings(query=TWO)) != job_key(ParseSettings(query="чайник"))
    assert job_key(ParseSettings(query=TWO)) == job_key(ParseSettings(query="Чайник\nТермокружка  "))


# --- collecting ---


def test_each_query_is_collected_and_marked():
    parser = FakeParser({"чайник": [product("1", 100)], "термокружка": [product("2", 200)]})
    found = asyncio.run(_search_all(parser, ParseSettings(query=TWO), Silent(), 0, 1))  # type: ignore[arg-type]
    assert parser.asked == ["чайник", "термокружка"]
    assert [(p.article, p.query) for p in found] == [("1", "чайник"), ("2", "термокружка")]


def test_a_product_found_by_two_queries_stays_twice():
    parser = FakeParser({"чайник": [product("1", 100)], "термокружка": [product("1", 100)]})
    found = asyncio.run(_search_all(parser, ParseSettings(query=TWO), Silent(), 0, 1))  # type: ignore[arg-type]
    assert [p.query for p in found] == ["чайник", "термокружка"]
    # Карточку и отзывы такого товара запрашиваем один раз.
    assert [p.query for p in _once_per_product(found)] == ["чайник"]


def test_the_twin_gets_the_card_of_the_first():
    found = [product("1", 100, "чайник"), product("1", None, "термокружка")]
    found[0].rating, found[0].brand, found[0].reviews_count = 4.8, "Bosch", 120
    _share_details(found)
    twin = found[1]
    assert (twin.rating, twin.brand, twin.reviews_count, twin.price) == (4.8, "Bosch", 120, 100)
    assert twin.query == "термокружка"  # запрос у близнеца свой


# --- what it means for the report ---


def test_rivals_come_from_the_same_query():
    products = [product("100", 500, "чайник", brand="Моя Марка"), product("1", 400, "чайник"),
                product("2", 300, "термокружка")]
    spots = compare_with_rivals(products, parse_own("", "Моя Марка")).spots
    assert len(spots) == 1
    assert [r.article for r in spots[0].rivals] == ["1"]  # товар из другого запроса не конкурент
    assert spots[0].rank == 2 and spots[0].total == 2


def test_the_file_is_named_after_the_first_query():
    name = build_file_name(ParseSettings(query=TWO), datetime(2026, 10, 22, 9, 0))
    assert name == "WB+Ozon_чайник и ещё 1_2026-10-22_09-00-00.xlsx"


@pytest.fixture
def two_queries() -> list[Product]:
    return [product("1", 100, "чайник"), product("2", 200, "чайник"), product("3", 300, "термокружка")]


def test_the_report_shows_which_query_found_what(tmp_path, two_queries):
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query=TWO), two_queries, [], datetime(2026, 10, 22, 9, 0))
    sheet = load_workbook(path)["Товары"]
    headers = [c.value for c in sheet[1]]
    assert headers[:3] == ["Площадка", "Артикул", "Запрос"]
    assert [sheet.cell(row=r, column=3).value for r in (2, 3, 4)] == ["чайник", "чайник", "термокружка"]

    summary = [row[0] for row in load_workbook(path)["Сводка"].iter_rows(values_only=True)]
    assert "Итоги по запросам" in summary and "Поисковые запросы" in summary


def test_a_single_query_leaves_no_empty_column(tmp_path, two_queries):
    path = tmp_path / "report.xlsx"
    one = [p for p in two_queries if p.query == "чайник"]
    export_to_excel(path, ParseSettings(query="чайник"), one, [], datetime(2026, 10, 22, 9, 0))
    workbook = load_workbook(path)
    assert "Запрос" not in [c.value for c in workbook["Товары"][1]]
    summary = [row[0] for row in workbook["Сводка"].iter_rows(values_only=True)]
    assert "Итоги по запросам" not in summary and "Поисковый запрос" in summary
