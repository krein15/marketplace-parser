"""Excel report and settings."""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import load_workbook

from mpparser.competitors import compare_with_rivals, parse_own
from mpparser.export.excel import build_file_name, export_to_excel, price_buckets
from mpparser.models import Product, Review
from mpparser.monitoring import CHEAPER, GONE, Change, Comparison, Dynamics, Positions
from mpparser.settings import InputMode, ParseSettings


@pytest.fixture
def products() -> list[Product]:
    return [
        Product(marketplace="Wildberries", article="1", name="Наушники WB", brand="Lirano", seller="Магазин",
                price=556.0, price_old=803.0, discount=31, rating=4.9, reviews_count=2569, position=1,
                url="https://www.wildberries.ru/catalog/1/detail.aspx", image="https://basket-01.wbbasket.ru/1.webp",
                category="Наушники", parsed_at=datetime(2026, 9, 18, 10, 30)),
        Product(marketplace="Ozon", article="2", name="Наушники Ozon", brand="JBL", seller="Продавец",
                price=3436.0, price_card=3100.0, price_old=5999.0, discount=42, rating=4.7, reviews_count=150,
                stock=287, position=1, url="https://www.ozon.ru/product/2/", category="Наушники"),
    ]


@pytest.fixture
def reviews() -> list[Review]:
    return [
        Review(marketplace="Wildberries", article="1", product_name="Наушники WB", date=datetime(2026, 9, 17, 20, 43),
               rating=5, author="Сержиккк", text="=) отличные наушники\nбас мощный", pros="цена", cons="",
               photos=2, likes=3, seller_answer="Спасибо!"),
        Review(marketplace="Ozon", article="2", product_name="Наушники Ozon", date=datetime(2026, 9, 13, 22, 43),
               rating=4, author="Ариша", text="Хорошие\x07наушники", variant="белый"),
    ]


def export(tmp_path, products, reviews, **overrides):
    settings = ParseSettings(query="наушники", **overrides)
    path = tmp_path / build_file_name(settings, datetime(2026, 9, 18, 10, 0))
    export_to_excel(path, settings, products, reviews, datetime(2026, 9, 18, 10, 0))
    return path


def test_file_name_contains_marketplaces_query_and_time():
    settings = ParseSettings(query="беспроводные наушники")
    assert build_file_name(settings, datetime(2026, 9, 18, 10, 0)) == \
        "WB+Ozon_беспроводные наушники_2026-09-18_10-00-00.xlsx"


def test_file_name_of_an_article_run_and_forbidden_characters():
    assert "артикулы" in build_file_name(ParseSettings(mode=InputMode.IDS), datetime(2026, 9, 18, 10, 0))
    name = build_file_name(ParseSettings(query='чехол/кейс: "про"'), datetime(2026, 9, 18, 10, 0))
    assert not set(name) & set('<>:"/\\|?*')


def test_sheets_and_headers(tmp_path, products, reviews):
    workbook = load_workbook(export(tmp_path, products, reviews))
    assert workbook.sheetnames == ["Сводка", "Товары", "Отзывы"]
    headers = [cell.value for cell in workbook["Товары"][1]]
    assert headers[:2] == ["Площадка", "Артикул"]
    assert "Цена, ₽" in headers
    assert workbook["Товары"].max_row == 3
    assert workbook["Товары"].freeze_panes == "C2"
    assert workbook["Товары"].tables


def test_reviews_sheet_is_skipped_when_reviews_are_off(tmp_path, products, reviews):
    workbook = load_workbook(export(tmp_path, products, reviews, collect_reviews=False))
    assert "Отзывы" not in workbook.sheetnames


def test_only_selected_columns_are_exported(tmp_path, products, reviews):
    workbook = load_workbook(export(tmp_path, products, reviews, product_fields=["price", "rating"]))
    assert [cell.value for cell in workbook["Товары"][1]] == ["Площадка", "Артикул", "Цена, ₽", "Рейтинг"]


def test_review_text_starting_with_equals_is_not_a_formula(tmp_path, products, reviews):
    sheet = load_workbook(export(tmp_path, products, reviews))["Отзывы"]
    text_column = [cell.value for cell in sheet[1]].index("Текст отзыва") + 1
    cell = sheet.cell(row=2, column=text_column)
    assert cell.data_type == "s"
    assert cell.value.startswith("=)")


def test_illegal_characters_are_stripped(tmp_path, products, reviews):
    sheet = load_workbook(export(tmp_path, products, reviews))["Отзывы"]
    text_column = [cell.value for cell in sheet[1]].index("Текст отзыва") + 1
    assert sheet.cell(row=3, column=text_column).value == "Хорошиенаушники"


def test_links_are_hyperlinks(tmp_path, products, reviews):
    sheet = load_workbook(export(tmp_path, products, reviews))["Товары"]
    url_column = [cell.value for cell in sheet[1]].index("Ссылка") + 1
    cell = sheet.cell(row=2, column=url_column)
    assert cell.hyperlink.target == "https://www.wildberries.ru/catalog/1/detail.aspx"
    assert cell.value == "Открыть"


def test_summary_statistics(tmp_path, products, reviews):
    sheet = load_workbook(export(tmp_path, products, reviews))["Сводка"]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    flat = {str(row[0]): row for row in values}
    assert flat["Собрано товаров"][1] == 2
    assert flat["Собрано отзывов"][1] == 2
    wb_row = next(row for row in values if row[0] == "Wildberries")
    assert wb_row[1] == 1  # products
    assert wb_row[2] == 556.0  # min price
    ozon_row = next(row for row in values if row[0] == "Ozon")
    assert ozon_row[3] == 3436.0  # average price


def test_export_without_reviews_still_writes_products(tmp_path, products):
    workbook = load_workbook(export(tmp_path, products, []))
    assert workbook["Отзывы"].max_row == 1  # header only


@pytest.mark.parametrize("prices", [[100.0], [100.0, 100.0, 100.0]])
def test_price_buckets_survive_degenerate_input(prices):
    buckets = price_buckets(prices)
    assert len(buckets) == 8
    assert buckets[0][0] <= min(prices)
    assert buckets[-1][1] is None


def test_price_buckets_cover_the_range():
    buckets = price_buckets([500.0, 1500.0, 9000.0, 25000.0])
    assert buckets[0][0] == 0
    assert all(low < high for low, high in buckets[:-1])


def test_changes_and_dynamics_sheets(tmp_path, products, reviews):
    comparison = Comparison(previous_at=datetime(2026, 9, 17, 9, 0), marketplaces=["Wildberries"], changes=[
        Change(CHEAPER, "Wildberries", "1", "Наушники WB", "https://wb/1", 600.0, 556.0),
        Change(GONE, "Wildberries", "7", "Старые наушники", "https://wb/7", 900.0, None),
    ])
    dynamics = Dynamics([datetime(2026, 9, 17, 9, 0), datetime(2026, 9, 18, 10, 0)],
                        [(products[0], [600.0, 556.0])])
    settings = ParseSettings(query="наушники", task_name="Наушники ежедневно")
    path = tmp_path / build_file_name(settings, datetime(2026, 9, 18, 10, 0))
    assert "Наушники ежедневно" in path.name
    export_to_excel(path, settings, products, reviews, datetime(2026, 9, 18, 10, 0), comparison, dynamics)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Сводка", "Изменения", "Товары", "Отзывы", "Динамика цен"]
    changes = workbook["Изменения"]
    assert [c.value for c in changes[4]][:3] == ["Статус", "Площадка", "Артикул"]
    assert [c.value for c in changes[5]][:8] == [CHEAPER, "Wildberries", "1", "Наушники WB", 600, 556, -44, -7.3]
    assert changes["I5"].hyperlink.target == "https://wb/1"
    assert changes["A6"].value == "Выпал из выдачи" and changes["F6"].value is None
    assert "рекламных мест" in changes["A3"].value
    history = workbook["Динамика цен"]
    assert [c.value for c in history[4]][:5] == ["Площадка", "Артикул", "Название", "17.09 09:00", "18.09 10:00"]
    assert [c.value for c in history[5]][3:5] == [600, 556]
    summary = [row[0] for row in workbook["Сводка"].iter_rows(values_only=True)]
    assert "Изменения с прошлого запуска" in summary and "Подешевели" in summary and "Задание" in summary
    assert "Выпали из выдачи" in summary


def test_positions_sheet(tmp_path, products, reviews):
    wb_product, ozon_product = products
    wb_product.position, ozon_product.position = 1, 12
    positions = Positions([datetime(2026, 9, 17, 9, 0), datetime(2026, 9, 18, 10, 0)],
                          [(ozon_product, [5, 12]), (wb_product, [3, 1])])
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query="наушники"), products, reviews, datetime(2026, 9, 18, 10, 0),
                    positions=positions)
    sheet = load_workbook(path)["Позиции"]
    assert [c.value for c in sheet[4]][:5] == ["Площадка", "Артикул", "Название", "Место сейчас", "Движение"]
    assert [c.value for c in sheet[5]][:5] == ["Ozon", "2", "Наушники Ozon", 12, -7]  # опустился на 7 мест
    assert [c.value for c in sheet[6]][:5] == ["Wildberries", "1", "Наушники WB", 1, 2]
    assert [c.value for c in sheet[5]][5:7] == [5, 12]


def test_positions_sheet_is_skipped_without_places(tmp_path, products, reviews):
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query="q"), products, reviews, datetime(2026, 9, 18),
                    positions=Positions([datetime(2026, 9, 18)], []))
    assert "Позиции" not in load_workbook(path).sheetnames


def test_first_run_changes_sheet_explains_itself(tmp_path, products, reviews):
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query="q"), products, reviews, datetime(2026, 9, 18), Comparison(None))
    workbook = load_workbook(path)
    assert "первый запуск" in workbook["Изменения"]["A2"].value
    assert "Динамика цен" not in workbook.sheetnames


def test_competitor_sheets(tmp_path, products, reviews):
    rivals = [
        Product(marketplace="Wildberries", article="7", name="Наушники дешевле", brand="Чужой", seller="Другой",
                price=499.0, rating=4.8, reviews_count=800, position=2, url="https://wb/7"),
        Product(marketplace="Wildberries", article="8", name="Наушники дороже", price=900.0, rating=4.0,
                url="https://wb/8"),
    ]
    collected = products + rivals
    competition = compare_with_rivals(collected, parse_own("1", ""))
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query="наушники"), collected, reviews, datetime(2026, 9, 18, 10, 0),
                    competition=competition)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Сводка", "Конкуренты", "Кто дешевле", "Товары", "Отзывы"]

    sheet = workbook["Конкуренты"]
    assert [c.value for c in sheet[4]][:6] == ["Площадка", "Артикул", "Мой товар", "Моя цена", "Место по цене",
                                               "Самая низкая цена"]
    assert [c.value for c in sheet[5]][:8] == ["Wildberries", "1", "Наушники WB", 556, "2 из 3", 499, 57, 11.4]
    assert [c.value for c in sheet[5]][8:10] == [699.5, -20.5]
    assert sheet["O5"].hyperlink.target == products[0].url

    cheaper = workbook["Кто дешевле"]
    assert [c.value for c in cheaper[4]][:4] == ["Мой артикул", "Площадка", "Конкурент", "Артикул"]
    assert [c.value for c in cheaper[5]][:9] == ["1", "Wildberries", "Наушники дешевле", "7", "Чужой", "Другой",
                                                 499, -57, -10.3]
    assert cheaper["M5"].hyperlink.target == "https://wb/7"
    summary = [row[0] for row in workbook["Сводка"].iter_rows(values_only=True)]
    assert "Мои товары против выдачи" in summary and "Дешевле меня" in summary


def test_the_cheaper_sheet_appears_only_when_somebody_undercuts_me(tmp_path, products, reviews):
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query="q"), products, reviews, datetime(2026, 9, 18),
                    competition=compare_with_rivals(products, parse_own("1", "")))
    names = load_workbook(path).sheetnames
    assert "Конкуренты" in names and "Кто дешевле" not in names


def test_competitor_sheets_are_skipped_without_my_products(tmp_path, products, reviews):
    path = tmp_path / "report.xlsx"
    export_to_excel(path, ParseSettings(query="q"), products, reviews, datetime(2026, 9, 18),
                    competition=compare_with_rivals(products, parse_own("", "")))
    assert "Конкуренты" not in load_workbook(path).sheetnames
