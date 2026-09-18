"""Excel report: "Товары", "Отзывы" and "Сводка" sheets."""

from __future__ import annotations

import math
import re
import statistics
from collections import Counter, defaultdict
from collections.abc import Iterable, Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from ..fields import PRODUCT_FIELDS, REVIEW_FIELDS, Field, resolve
from ..models import Product, Review
from ..regions import WB_REGIONS
from ..settings import MARKETPLACES, SORT_TITLES, InputMode, ParseSettings

HEADER_FILL = PatternFill("solid", fgColor="2B2D42")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="2B2D42")
SECTION_FONT = Font(bold=True, size=12, color="2B2D42")
MUTED_FONT = Font(color="6B7280")
MARKETPLACE_FONTS = {"Wildberries": Font(bold=True, color="A20D8A"), "Ozon": Font(bold=True, color="005BFF")}
TABLE_STYLE = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)

NUMBER_FORMATS = {
    "money": '#,##0 "₽"',
    "percent": '0"%"',
    "rating": "0.0",
    "int": "#,##0",
    "datetime": "DD.MM.YYYY HH:MM",
}
URL_LABELS = {"image": "Фото", "url": "Открыть"}
MAX_CELL_TEXT = 32_000
MAX_ROW_HEIGHT = 150
LINE_HEIGHT = 15


def build_file_name(settings: ParseSettings, started_at: datetime) -> str:
    marketplaces = "+".join("WB" if key == "wb" else MARKETPLACES[key] for key in settings.marketplaces)
    subject = settings.query.strip() if settings.mode == InputMode.QUERY else "артикулы"
    subject = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", subject).strip()[:40].strip() or "сбор"
    return f"{marketplaces}_{subject}_{started_at:%Y-%m-%d_%H-%M-%S}.xlsx"


def export_to_excel(
    path: Path,
    settings: ParseSettings,
    products: Sequence[Product],
    reviews: Sequence[Review],
    started_at: datetime,
) -> Path:
    workbook = Workbook()
    products_sheet = workbook.active
    products_sheet.title = "Товары"
    _write_table(products_sheet, "Products", resolve(PRODUCT_FIELDS, settings.product_fields), products, "C2")
    if settings.collect_reviews:
        _write_table(workbook.create_sheet("Отзывы"), "Reviews", resolve(REVIEW_FIELDS, settings.review_fields),
                     reviews, "C2")
    _write_summary(workbook.create_sheet("Сводка", 0), settings, products, reviews, started_at)
    workbook.active = 0

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


# --- data sheets ---


def _cell_value(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_TEXT]
    return value


def _write_table(sheet: Worksheet, name: str, columns: list[Field], rows: Sequence[Any], freeze: str) -> None:
    for col, spec in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=col, value=spec.title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col)].width = spec.width
    sheet.row_dimensions[1].height = 32

    for row_index, item in enumerate(rows, 2):
        lines = 1
        for col, spec in enumerate(columns, 1):
            value = _cell_value(getattr(item, spec.key))
            cell = sheet.cell(row=row_index, column=col)
            if spec.kind == "url":
                if value:
                    cell.value, cell.hyperlink, cell.style = URL_LABELS.get(spec.key, "Открыть"), value, "Hyperlink"
                continue
            cell.value = value if value != "" else None
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # review text like "=)" must not become a formula
            if spec.kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[spec.kind]
            if spec.kind == "wrap":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(value, str):
                    lines = max(lines, _estimate_lines(value, spec.width))
            else:
                cell.alignment = Alignment(vertical="top")
            if spec.key == "marketplace" and value in MARKETPLACE_FONTS:
                cell.font = MARKETPLACE_FONTS[value]
        if lines > 1:
            sheet.row_dimensions[row_index].height = min(lines * LINE_HEIGHT, MAX_ROW_HEIGHT)

    last_row = max(len(rows) + 1, 2)  # an Excel table needs at least one data row
    table = Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{last_row}")
    table.tableStyleInfo = TABLE_STYLE
    sheet.add_table(table)
    sheet.freeze_panes = freeze


def _estimate_lines(text: str, width: int) -> int:
    """Excel does not auto-fit row heights of generated files, so estimate the wrapped line count."""
    chars_per_line = max(int(width * 1.15), 1)
    return sum(max(math.ceil(len(line) / chars_per_line), 1) for line in text.splitlines() or [""])


# --- summary sheet ---


def _write_summary(
    sheet: Worksheet,
    settings: ParseSettings,
    products: Sequence[Product],
    reviews: Sequence[Review],
    started_at: datetime,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "Сбор данных с маркетплейсов"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Сформировано {started_at:%d.%m.%Y в %H:%M}"
    sheet["A2"].font = MUTED_FONT
    for col, width in zip("ABCDEFGHIJ", (24, 12, 13, 15, 13, 13, 14, 15, 16, 15), strict=True):
        sheet.column_dimensions[col].width = width

    if settings.mode == InputMode.QUERY:
        subject = ("Поисковый запрос", settings.query.strip())
    else:
        subject = ("Артикулы", "список артикулов и ссылок")
    params = [
        ("Площадки", ", ".join(MARKETPLACES[k] for k in settings.marketplaces)),
        subject,
        ("Сортировка", SORT_TITLES[settings.sort] if settings.mode == InputMode.QUERY else "—"),
        ("Регион WB", settings.region if settings.region in WB_REGIONS else "—"),
        ("Собрано товаров", len(products)),
        ("Собрано отзывов", len(reviews) if settings.collect_reviews else "не собирались"),
    ]
    row = _section(sheet, 4, "Параметры сбора")
    for label, value in params:
        sheet.cell(row=row, column=1, value=label).font = MUTED_FONT
        sheet.cell(row=row, column=2, value=value).alignment = Alignment(horizontal="left")
        row += 1

    row = _section(sheet, row + 1, "Итоги по площадкам")
    headers = ["Площадка", "Товаров", "Мин. цена", "Средняя цена", "Медиана", "Макс. цена",
               "Средняя скидка", "Средний рейтинг", "Отзывов на сайте", "Собрано отзывов"]
    formats = [None, "#,##0", *[NUMBER_FORMATS["money"]] * 4, NUMBER_FORMATS["percent"], "0.00", "#,##0", "#,##0"]
    stats_rows = []
    for key in settings.marketplaces:
        title = MARKETPLACES[key]
        items = [p for p in products if p.marketplace == title]
        prices = [p.price for p in items if p.price]
        stats_rows.append([
            title,
            len(items),
            min(prices, default=None),
            _mean(prices),
            statistics.median(prices) if prices else None,
            max(prices, default=None),
            _mean(p.discount for p in items if p.discount),
            _mean(p.rating for p in items if p.rating),
            sum(p.reviews_count or 0 for p in items),
            sum(1 for r in reviews if r.marketplace == title),
        ])
    row = _small_table(sheet, row, 1, headers, stats_rows, formats)

    top_row = _section(sheet, row + 1, "Топ-10 брендов")
    _section(sheet, row + 1, "Топ-10 продавцов", column=6)
    brand_rows = _top_groups(products, "brand")
    seller_rows = _top_groups(products, "seller")
    group_formats = [None, "#,##0", NUMBER_FORMATS["money"], "0.00"]
    end_brands = _small_table(sheet, top_row, 1, ["Бренд", "Товаров", "Средняя цена", "Рейтинг"], brand_rows,
                              group_formats)
    end_sellers = _small_table(sheet, top_row, 6, ["Продавец", "Товаров", "Средняя цена", "Рейтинг"], seller_rows,
                               group_formats)
    row = max(end_brands, end_sellers)

    prices_by_mp = {MARKETPLACES[k]: [p.price for p in products if p.marketplace == MARKETPLACES[k] and p.price]
                    for k in settings.marketplaces}
    if sum(len(v) for v in prices_by_mp.values()) >= 5:
        _price_chart(sheet, row + 1, prices_by_mp)


def _section(sheet: Worksheet, row: int, title: str, column: int = 1) -> int:
    sheet.cell(row=row, column=column, value=title).font = SECTION_FONT
    return row + 1


def _small_table(
    sheet: Worksheet, row: int, column: int, headers: list[str], rows: list[list[Any]], formats: list[str | None]
) -> int:
    for offset, header in enumerate(headers):
        cell = sheet.cell(row=row, column=column + offset, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 30
    if not rows:
        sheet.cell(row=row + 1, column=column, value="нет данных").font = MUTED_FONT
        return row + 2
    for r_offset, values in enumerate(rows, 1):
        for offset, value in enumerate(values):
            cell = sheet.cell(row=row + r_offset, column=column + offset, value=_cell_value(value))
            if formats[offset]:
                cell.number_format = formats[offset]
            if offset == 0 and value in MARKETPLACE_FONTS:
                cell.font = MARKETPLACE_FONTS[value]
    return row + len(rows) + 1


def _mean(values: Iterable[float | int | None]) -> float | None:
    items = [v for v in values if v is not None]
    return round(statistics.fmean(items), 2) if items else None


def _top_groups(products: Sequence[Product], attribute: str, limit: int = 10) -> list[list[Any]]:
    groups: dict[str, list[Product]] = defaultdict(list)
    for product in products:
        if name := getattr(product, attribute):
            groups[name].append(product)
    counts = Counter({name: len(items) for name, items in groups.items()})
    return [
        [name, count, _mean(p.price for p in groups[name]), _mean(p.rating for p in groups[name])]
        for name, count in counts.most_common(limit)
    ]


def _nice_step(raw: float) -> float:
    exponent = 10 ** math.floor(math.log10(raw)) if raw > 0 else 1
    for factor in (1, 2, 2.5, 5, 10):
        if raw <= factor * exponent:
            return factor * exponent
    return 10 * exponent


def price_buckets(prices: Sequence[float], count: int = 8) -> list[tuple[float, float | None]]:
    """Equal-width price ranges with round bounds; the last one is open-ended to absorb outliers."""
    ordered = sorted(prices)
    low, high = ordered[0], ordered[max(int(len(ordered) * 0.9) - 1, 0)]
    step = _nice_step(max((high - low) / count, 1))
    start = math.floor(low / step) * step
    edges = [(start + i * step, start + (i + 1) * step) for i in range(count - 1)]
    return [*edges, (start + (count - 1) * step, None)]


def _price_chart(sheet: Worksheet, row: int, prices_by_mp: dict[str, list[float]]) -> None:
    row = _section(sheet, row, "Распределение цен")
    buckets = price_buckets([p for prices in prices_by_mp.values() for p in prices])
    names = [name for name, prices in prices_by_mp.items() if prices]
    table_rows = []
    for low, high in buckets:
        label = f"{low:,.0f}–{high:,.0f} ₽" if high is not None else f"от {low:,.0f} ₽"
        counts = [sum(1 for p in prices_by_mp[n] if p >= low and (high is None or p < high)) for n in names]
        table_rows.append([label.replace(",", " "), *counts])
    end = _small_table(sheet, row, 1, ["Диапазон цен", *names], table_rows, [None, *["#,##0"] * len(names)])

    chart = BarChart()
    chart.type, chart.grouping = "col", "clustered"
    chart.title = "Количество товаров по ценовым диапазонам"
    chart.y_axis.title = "Товаров"
    chart.height, chart.width = 8, 22
    data = Reference(sheet, min_col=2, max_col=1 + len(names), min_row=row, max_row=end - 1)
    labels = Reference(sheet, min_col=1, min_row=row + 1, max_row=end - 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    colors = {"Wildberries": "A20D8A", "Ozon": "005BFF"}
    for series, name in zip(chart.series, names, strict=True):
        series.graphicalProperties.solidFill = colors.get(name, "2B2D42")
        series.graphicalProperties.line.solidFill = colors.get(name, "2B2D42")
    sheet.add_chart(chart, f"D{row}")
